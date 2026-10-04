"""The three files an assistant makes here rather than in its own container.

The assistant's container runs no Python, so a skill that told it to build a
workbook with openpyxl, a deck with python-pptx or a PDF with reportlab named
libraries it could not import. Each of those files is made here instead:

  - `convert_md_to_pptx`: markdown to editable slides, through pandoc.
  - `convert_html_to_pdf`: a page to PDF through headless Chromium. LibreOffice
    also takes HTML, and it drops CSS grid, flexbox and background colours, so a
    landing page came out as a column of unstyled boxes.
  - `create_xlsx`: a workbook from rows, through openpyxl, with formulas, a bold
    frozen header, number formats and several sheets.

pandoc and Chromium are faked as in test_server.py; openpyxl runs for real,
because what is asserted is the workbook it writes.

    python -m pytest tests/ -q
"""
from __future__ import annotations

import base64
import datetime
import io
import os
import subprocess

import pytest
from openpyxl import load_workbook

import server

AGENT = "agent-7"


class FakeBinaries:
    """Records argv and writes the file pandoc or Chromium would have written."""

    def __init__(self):
        self.calls: list[list[str]] = []
        self.payload = b"%PDF-1.7 fake"
        self.seen_html: str | None = None

    def __call__(self, argv, **kwargs):
        argv = list(argv)
        self.calls.append(argv)
        if argv[0] == server.PANDOC:
            out = argv[argv.index("-o") + 1]
        else:
            out = next(a.split("=", 1)[1] for a in argv if a.startswith("--print-to-pdf="))
            src = argv[-1]
            assert src.startswith("file://")
            with open(src[len("file://"):], encoding="utf-8") as f:
                self.seen_html = f.read()
        with open(out, "wb") as f:
            f.write(self.payload)
        return subprocess.CompletedProcess(argv, 0, b"", b"")

    def argv_for(self, binary: str) -> list[str]:
        for argv in self.calls:
            if argv[0] == binary:
                return argv
        raise AssertionError(f"{binary} was never invoked; calls={self.calls}")


class FakeResponse:
    status = 200

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def read(self):
        return b""


class FakeBroker:
    def __init__(self):
        self.requests = []

    def __call__(self, req, timeout=None):
        self.requests.append(req)
        return FakeResponse()


@pytest.fixture
def binaries(monkeypatch):
    fake = FakeBinaries()
    monkeypatch.setattr(server.subprocess, "run", fake)
    return fake


@pytest.fixture
def broker(monkeypatch):
    monkeypatch.setenv("CERASE_CONTROL_PLANE_URL", "http://cerase-control-plane")
    monkeypatch.setenv("CERASE_INTERNAL_SECRET", "internal-secret")
    fake = FakeBroker()
    monkeypatch.setattr(server.urllib.request, "urlopen", fake)
    return fake


@pytest.fixture
def no_broker(monkeypatch):
    monkeypatch.delenv("CERASE_CONTROL_PLANE_URL", raising=False)
    monkeypatch.delenv("CERASE_INTERNAL_SECRET", raising=False)


def _b64(text: str) -> str:
    return base64.b64encode(text.encode()).decode()


# ─── Markdown to PPTX ────────────────────────────────────────────────────

def test_markdown_becomes_a_pptx_through_pandoc(binaries, broker):
    result = server.convert_md_to_pptx(
        input_b64=_b64("# Cover\n\n## Slide\n\n- one\n"),
        agent_id=AGENT,
        output_filename="q3.pptx",
    )
    argv = binaries.argv_for(server.PANDOC)
    assert argv[argv.index("-o") + 1].endswith(".pptx")
    assert result == {"path": "outputs/q3.pptx", "filename": "q3.pptx", "size_bytes": len(binaries.payload)}


def test_a_pptx_template_reaches_pandoc(binaries, broker):
    server.convert_md_to_pptx(
        input_b64=_b64("# Cover\n"),
        reference_doc_b64=base64.b64encode(b"template pptx").decode(),
        agent_id=AGENT,
    )
    ref = [a for a in binaries.argv_for(server.PANDOC) if a.startswith("--reference-doc=")]
    assert len(ref) == 1 and ref[0].endswith("reference.pptx")


# ─── HTML to PDF ─────────────────────────────────────────────────────────

PAGE = "<!doctype html><html><head><style>.grid{display:grid}</style></head><body><h1>Acme</h1></body></html>"


def test_html_becomes_a_pdf_through_chromium_and_not_libreoffice(binaries, broker):
    result = server.convert_html_to_pdf(input_b64=_b64(PAGE), agent_id=AGENT, output_filename="landing.pdf")
    argv = binaries.argv_for(server.CHROMIUM)
    assert "--headless" in argv
    assert any(a.startswith("--print-to-pdf=") for a in argv)
    assert all(c[0] != server.SOFFICE for c in binaries.calls)
    assert result["path"] == "outputs/landing.pdf"


def test_the_page_size_is_set_before_the_pages_own_css(binaries, broker):
    server.convert_html_to_pdf(input_b64=_b64(PAGE), agent_id=AGENT, paper="letter", orientation="landscape")
    html = binaries.seen_html
    injected = html.index("size: letter landscape")
    # Before the page's own <style>, so a page that sets its own @page wins.
    assert injected < html.index(".grid{display:grid}")
    # Backgrounds print: a hero band in the brand colour is the page.
    assert "print-color-adjust: exact" in html


def test_a_page_with_no_head_still_gets_a_size(binaries, broker):
    server.convert_html_to_pdf(input_b64=_b64("<p>bare</p>"), agent_id=AGENT)
    assert "size: A4 portrait" in binaries.seen_html
    assert "<p>bare</p>" in binaries.seen_html


def test_html_to_pdf_refuses_an_unknown_paper_or_orientation(binaries, broker):
    with pytest.raises(ValueError, match="paper"):
        server.convert_html_to_pdf(input_b64=_b64(PAGE), agent_id=AGENT, paper="A0")
    with pytest.raises(ValueError, match="orientation"):
        server.convert_html_to_pdf(input_b64=_b64(PAGE), agent_id=AGENT, orientation="sideways")


def test_html_to_pdf_names_the_output_after_the_input_when_no_name_is_given(binaries, broker):
    result = server.convert_html_to_pdf(input_b64=_b64(PAGE), agent_id=AGENT)
    assert result["filename"].endswith(".pdf")


# ─── Spreadsheets ────────────────────────────────────────────────────────

def _workbook(no_broker_result: dict):
    return load_workbook(io.BytesIO(base64.b64decode(no_broker_result["contents_base64"])))


def test_a_workbook_keeps_formulas_numbers_and_a_frozen_bold_header(no_broker):
    result = server.create_xlsx(
        sheets=[{
            "name": "Sales",
            "rows": [["Region", "Q1", "Q2", "Total"], ["North", 10, 20.5, "=SUM(B2:C2)"], ["South", 5, 7, "=SUM(B3:C3)"]],
        }],
        output_filename="sales.xlsx",
    )
    ws = _workbook(result)["Sales"]
    assert ws["D2"].value == "=SUM(B2:C2)"
    assert ws["B2"].value == 10 and ws["C2"].value == 20.5
    assert ws["A1"].font.bold is True
    assert ws.freeze_panes == "A2"
    assert result["filename"] == "sales.xlsx"


def test_number_formats_and_widths_apply_to_the_named_columns(no_broker):
    result = server.create_xlsx(
        sheets=[{
            "name": "Costs",
            "rows": [["Item", "Amount", "Share"], ["Rent", 1200, 0.42]],
            "number_formats": {"B": '#,##0.00 "€"', "C": "0.0%"},
            "column_widths": {"A": 30},
        }],
    )
    ws = _workbook(result)["Costs"]
    assert ws["B2"].number_format == '#,##0.00 "€"'
    assert ws["C2"].number_format == "0.0%"
    assert ws.column_dimensions["A"].width == 30
    # The header keeps its own format: it is text.
    assert ws["B1"].number_format == "General"


def test_an_iso_date_is_stored_as_a_date(no_broker):
    result = server.create_xlsx(sheets=[{"name": "Log", "rows": [["Day"], ["2026-10-05"]]}])
    cell = _workbook(result)["Log"]["A2"]
    assert cell.value == datetime.datetime(2026, 10, 5)
    assert cell.number_format == "yyyy-mm-dd"


def test_several_sheets_keep_their_order(no_broker):
    result = server.create_xlsx(sheets=[
        {"name": "January", "rows": [["a"], [1]]},
        {"name": "February", "rows": [["a"], [2]]},
    ])
    assert _workbook(result).sheetnames == ["January", "February"]


def test_a_sheet_without_a_header_is_not_frozen(no_broker):
    result = server.create_xlsx(sheets=[{"name": "Raw", "rows": [[1, 2]], "header_rows": 0}])
    ws = _workbook(result)["Raw"]
    assert ws.freeze_panes is None
    assert ws["A1"].font.bold is not True


def test_a_workbook_goes_to_the_workspace_as_a_path(broker):
    result = server.create_xlsx(sheets=[{"name": "S", "rows": [["a"], [1]]}], agent_id=AGENT, output_filename="s.xlsx")
    assert result["path"] == "outputs/s.xlsx"
    assert "contents_base64" not in result
    assert broker.requests[-1].get_method() == "PUT"


@pytest.mark.parametrize("sheets, message", [
    ([], "at least one sheet"),
    ([{"name": "A", "rows": []}], "no rows"),
    ([{"name": "A", "rows": [[1]]}, {"name": "a", "rows": [[1]]}], "twice"),
    ([{"name": "Bad/Name", "rows": [[1]]}], "cannot contain"),
    ([{"name": "x" * 32, "rows": [[1]]}], "31"),
])
def test_a_workbook_that_excel_would_refuse_is_refused_here(no_broker, sheets, message):
    with pytest.raises(ValueError, match=message):
        server.create_xlsx(sheets=sheets)


def test_the_output_name_gets_its_extension(no_broker):
    result = server.create_xlsx(sheets=[{"name": "S", "rows": [[1]]}], output_filename="report")
    assert result["filename"] == "report.xlsx"



def test_chromium_prints_no_header_or_footer_of_its_own(binaries, broker):
    server.convert_html_to_pdf(input_b64=_b64(PAGE), agent_id=AGENT)
    assert "--no-pdf-header-footer" in binaries.argv_for(server.CHROMIUM)
